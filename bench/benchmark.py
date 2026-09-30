"""
benchmark.py — Menjalankan SELURUH pengujian wajib Topik A dan menulis
hasilnya ke hasil/hasil_pengujian.xlsx + grafik PNG di hasil/grafik/.

Jalankan dari akar proyek:  python bench/benchmark.py

Pengujian yang dijalankan
  T1  Kebenaran dekripsi pada 11 masukan berbeda (termasuk PNG dan PDF)
  T2  Waktu enkripsi & dekripsi untuk 1 KB, 1 MB, 10 MB
  T3  Avalanche effect (1 bit plainteks diubah, dan 1 bit kunci diubah)
  T4  Entropi Shannon + histogram byte + uji chi-square keseragaman
  T5  Perbandingan AES-256-GCM vs ChaCha20-Poly1305 (+ AES-CBC & ECB sebagai
      PEMBANDING yang tidak dipakai sebagai fitur keamanan utama)
"""
import os
import secrets
import statistics
import sys
import time

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

AKAR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, AKAR)
from core import crypto_core as cc

from cryptography.hazmat.primitives import padding as sympad
from cryptography.hazmat.primitives.ciphers import Cipher, algorithms, modes

DIR_UJI = os.path.join(AKAR, "data", "uji")
DIR_HASIL = os.path.join(AKAR, "hasil")
DIR_GRAFIK = os.path.join(DIR_HASIL, "grafik")
os.makedirs(DIR_GRAFIK, exist_ok=True)

# Kata sandi benchmark diambil dari variabel lingkungan, atau dibangkitkan acak
# bila tidak diset. Sengaja TIDAK ditulis sebagai literal di kode sumber, sesuai
# ketentuan teknis tugas (tidak ada kata sandi di dalam repositori).
SANDI = os.environ.get("BENCH_SANDI") or secrets.token_urlsafe(24)
WARNA = {"AES-256-GCM": "#4f9cf9", "ChaCha20-Poly1305": "#f0883e",
         "AES-256-CBC": "#8b949e", "AES-256-ECB": "#f85149"}

plt.rcParams.update({"figure.dpi": 130, "font.size": 9,
                     "axes.grid": True, "grid.alpha": .25,
                     "axes.spines.top": False, "axes.spines.right": False})


def baca(nama):
    with open(os.path.join(DIR_UJI, nama), "rb") as f:
        return f.read()


def berkas_uji():
    return sorted(os.listdir(DIR_UJI))


# ===========================================================================
# T1 — Kebenaran dekripsi
# ===========================================================================
def t1_kebenaran():
    print("\n[T1] Uji kebenaran dekripsi")
    baris = []
    for nama in berkas_uji():
        data = baca(nama)
        for alg_id, alg in cc.ALG_NAMES.items():
            blob = cc.encrypt(data, SANDI, alg_id=alg_id)
            pulih = cc.decrypt(blob, SANDI)
            cocok = pulih == data

            # verifikasi penolakan
            try:
                cc.decrypt(blob, SANDI + "x"); tolak_sandi = False
            except cc.CryptoError:
                tolak_sandi = True
            rusak = bytearray(blob); rusak[-1] ^= 0x01
            try:
                cc.decrypt(bytes(rusak), SANDI); tolak_ubah = False
            except cc.CryptoError:
                tolak_ubah = True

            baris.append([nama, alg, len(data), len(blob), len(blob) - len(data),
                          "BERHASIL" if cocok else "GAGAL",
                          "DITOLAK" if tolak_sandi else "LOLOS(!)",
                          "DITOLAK" if tolak_ubah else "LOLOS(!)"])
    lulus = sum(1 for b in baris if b[5] == "BERHASIL" and b[6] == "DITOLAK" and b[7] == "DITOLAK")
    print(f"      {lulus}/{len(baris)} kasus lulus penuh")
    return baris


# ===========================================================================
# T2 — Waktu eksekusi
# ===========================================================================
def t2_waktu(ulang=7):
    print("\n[T2] Uji waktu enkripsi/dekripsi")
    kunci = os.urandom(32)
    baris, grafik = [], {}
    berkas = [("1 KB", "08_biner_1KB.bin"), ("1 MB", "09_biner_1MB.bin"),
              ("10 MB", "10_biner_10MB.bin")]

    for alg_id, alg in cc.ALG_NAMES.items():
        grafik[alg] = {"label": [], "enk": [], "dek": [], "tput": []}
        for label, nama in berkas:
            data = baca(nama)
            nonce = os.urandom(12)
            aead = cc._aead(alg_id, kunci)

            te, td = [], []
            for _ in range(ulang):
                t = time.perf_counter(); ct = aead.encrypt(nonce, data, b""); te.append(time.perf_counter() - t)
                t = time.perf_counter(); aead.decrypt(nonce, ct, b""); td.append(time.perf_counter() - t)
            me, md = statistics.median(te) * 1000, statistics.median(td) * 1000
            tput = len(data) / (me / 1000) / 1e6      # MB/s
            baris.append([alg, label, len(data), round(me, 3), round(md, 3),
                          round(statistics.stdev(te) * 1000, 3), round(tput, 1)])
            grafik[alg]["label"].append(label)
            grafik[alg]["enk"].append(me); grafik[alg]["dek"].append(md)
            grafik[alg]["tput"].append(tput)
            print(f"      {alg:20} {label:>6}  enk {me:8.3f} ms  dek {md:8.3f} ms  {tput:6.1f} MB/s")

    # biaya KDF diukur terpisah — ini konstanta per operasi, bukan per byte
    kdf_baris = []
    for kdf_id, kdf in cc.KDF_NAMES.items():
        try:
            t = time.perf_counter()
            cc.derive_key(SANDI, os.urandom(16), kdf_id)
            ms = (time.perf_counter() - t) * 1000
            kdf_baris.append([kdf, round(ms, 2), "tersedia"])
            print(f"      KDF {kdf:22} {ms:8.2f} ms")
        except cc.CryptoError as e:
            kdf_baris.append([kdf, None, f"tidak tersedia ({e})"])

    _grafik_waktu(grafik)
    return baris, kdf_baris


def _grafik_waktu(g):
    fig, ax = plt.subplots(1, 2, figsize=(9.5, 3.4))
    x = np.arange(3); w = 0.35
    for i, (alg, d) in enumerate(g.items()):
        ax[0].bar(x + (i - .5) * w, d["enk"], w, label=alg, color=WARNA[alg])
        ax[1].bar(x + (i - .5) * w, d["tput"], w, label=alg, color=WARNA[alg])
    ax[0].set_yscale("log"); ax[0].set_ylabel("waktu enkripsi (ms, skala log)")
    ax[1].set_ylabel("throughput (MB/s)")
    for a in ax:
        a.set_xticks(x); a.set_xticklabels(["1 KB", "1 MB", "10 MB"]); a.legend(fontsize=8)
    ax[0].set_title("Waktu enkripsi menurut ukuran berkas", fontsize=10)
    ax[1].set_title("Throughput", fontsize=10)
    fig.tight_layout(); fig.savefig(os.path.join(DIR_GRAFIK, "t2_waktu.png")); plt.close(fig)


# ===========================================================================
# T3 — Avalanche effect
# ===========================================================================
def _bit_beda(a: bytes, b: bytes) -> tuple[int, int]:
    n = min(len(a), len(b))
    x = np.frombuffer(a[:n], np.uint8) ^ np.frombuffer(b[:n], np.uint8)
    return int(np.unpackbits(x).sum()), n * 8


def _cbc(kunci, iv, data):
    p = sympad.PKCS7(128).padder()
    enc = Cipher(algorithms.AES(kunci), modes.CBC(iv)).encryptor()
    return enc.update(p.update(data) + p.finalize()) + enc.finalize()


def t3_avalanche(percobaan=100, n_byte=1024):
    print(f"\n[T3] Avalanche effect ({percobaan} percobaan × {n_byte} byte)")
    rng = np.random.default_rng(7)
    hasil, baris = {}, []

    skema = [(cc.ALG_AES_GCM, "AES-256-GCM"), (cc.ALG_CHACHA20, "ChaCha20-Poly1305"),
             (None, "AES-256-CBC")]

    for alg_id, alg in skema:
        for sumber in ("plainteks", "kunci"):
            persen_badan, persen_tag = [], []
            for _ in range(percobaan):
                kunci = bytes(rng.integers(0, 256, 32, dtype=np.uint8))
                nonce = bytes(rng.integers(0, 256, 16 if alg_id is None else 12, dtype=np.uint8))
                plain = bytes(rng.integers(0, 256, n_byte, dtype=np.uint8))

                def enk(k, p):
                    if alg_id is None:
                        return _cbc(k, nonce, p)
                    return cc._aead(alg_id, k).encrypt(nonce, p, b"")

                c1 = enk(kunci, plain)
                if sumber == "plainteks":
                    p2 = bytearray(plain); p2[rng.integers(0, n_byte)] ^= 1 << rng.integers(0, 8)
                    c2 = enk(kunci, bytes(p2))
                else:
                    k2 = bytearray(kunci); k2[rng.integers(0, 32)] ^= 1 << rng.integers(0, 8)
                    c2 = enk(bytes(k2), plain)

                if alg_id is None:
                    d, t = _bit_beda(c1, c2); persen_badan.append(d / t * 100)
                else:
                    d, t = _bit_beda(c1[:n_byte], c2[:n_byte]); persen_badan.append(d / t * 100)
                    dt, tt = _bit_beda(c1[n_byte:], c2[n_byte:]); persen_tag.append(dt / tt * 100)

            rb = statistics.mean(persen_badan)
            rt = statistics.mean(persen_tag) if persen_tag else None
            hasil[(alg, sumber)] = persen_badan
            baris.append([alg, f"1 bit {sumber} diubah", round(rb, 3),
                          round(statistics.stdev(persen_badan), 3),
                          round(rt, 3) if rt is not None else "—",
                          "ideal 50%"])
            print(f"      {alg:20} ubah {sumber:10} badan {rb:7.3f}%"
                  + (f"  tag {rt:6.2f}%" if rt is not None else ""))

    _grafik_avalanche(baris)
    return baris


def _grafik_avalanche(baris):
    fig, ax = plt.subplots(figsize=(8, 3.6))
    label = [f"{b[0]}\n{b[1].replace('1 bit ','').replace(' diubah','')}" for b in baris]
    badan = [b[2] for b in baris]
    tag = [b[4] if isinstance(b[4], (int, float)) else 0 for b in baris]
    x = np.arange(len(baris)); w = .38
    ax.bar(x - w / 2, badan, w, label="badan cipherteks", color="#4f9cf9")
    ax.bar(x + w / 2, tag, w, label="authentication tag", color="#3fb950")
    ax.axhline(50, ls="--", c="#f85149", lw=1, label="ideal 50%")
    ax.set_xticks(x); ax.set_xticklabels(label, fontsize=7.5)
    ax.set_ylabel("bit berubah (%)"); ax.set_ylim(0, 60); ax.legend(fontsize=8)
    ax.set_title("Avalanche effect — mode stream (GCM/ChaCha20) vs blok (CBC)", fontsize=10)
    fig.tight_layout(); fig.savefig(os.path.join(DIR_GRAFIK, "t3_avalanche.png")); plt.close(fig)


# ===========================================================================
# T4 — Entropi, histogram, chi-square
# ===========================================================================
def _entropi(data: bytes) -> float:
    c = np.bincount(np.frombuffer(data, np.uint8), minlength=256)
    p = c[c > 0] / len(data)
    return float(-(p * np.log2(p)).sum())


def _chi2(data: bytes) -> float:
    """Chi-square terhadap sebaran seragam (256 nilai, df=255).
    Nilai kritis 5% ≈ 293,25; di bawah itu = tidak bisa dibedakan dari acak."""
    c = np.bincount(np.frombuffer(data, np.uint8), minlength=256)
    h = len(data) / 256
    return float(((c - h) ** 2 / h).sum())


def t4_entropi():
    print("\n[T4] Entropi, histogram, chi-square")
    sasaran = ["02_teks_indonesia.txt", "06_citra.png", "07_dokumen.pdf",
               "09_biner_1MB.bin", "11_berpola_1MB.bin"]
    baris = []
    for nama in sasaran:
        data = baca(nama)
        if len(data) < 4096:
            data = data * (4096 // len(data) + 1)   # perbesar agar statistik stabil
        ep, xp = _entropi(data), _chi2(data)
        for alg_id, alg in cc.ALG_NAMES.items():
            ct = cc.encrypt(data, SANDI, alg_id=alg_id)[cc.HEADER_LEN:]
            baris.append([nama, alg, len(data), round(ep, 4), round(_entropi(ct), 4),
                          round(xp, 1), round(_chi2(ct), 1),
                          "seragam" if _chi2(ct) < 293.25 else "tidak seragam"])
        print(f"      {nama:24} plain {ep:5.3f} -> cipher {baris[-1][4]:5.3f} bit/byte")

    _grafik_histogram()
    _grafik_entropi(baris)
    return baris


def _grafik_histogram():
    data = baca("11_berpola_1MB.bin")
    ct = cc.encrypt(data, SANDI)[cc.HEADER_LEN:]
    citra = baca("06_citra.png") * 40
    fig, ax = plt.subplots(1, 3, figsize=(10.5, 2.9))
    for a, (d, judul, warna) in zip(ax, [
            (data, f"Plainteks berpola\nH={_entropi(data):.3f} bit/byte", "#8b949e"),
            (citra, f"Plainteks citra PNG\nH={_entropi(citra):.3f} bit/byte", "#f0883e"),
            (ct, f"Cipherteks AES-256-GCM\nH={_entropi(ct):.3f} bit/byte", "#4f9cf9")]):
        a.hist(np.frombuffer(d, np.uint8), bins=256, range=(0, 255), color=warna)
        a.set_title(judul, fontsize=9); a.set_xlabel("nilai byte"); a.set_xlim(0, 255)
    ax[0].set_ylabel("frekuensi")
    fig.tight_layout(); fig.savefig(os.path.join(DIR_GRAFIK, "t4_histogram.png")); plt.close(fig)


def _grafik_entropi(baris):
    nama = sorted({b[0] for b in baris})
    fig, ax = plt.subplots(figsize=(8, 3.2))
    x = np.arange(len(nama)); w = .27
    ax.bar(x - w, [next(b[3] for b in baris if b[0] == n) for n in nama], w,
           label="plainteks", color="#8b949e")
    for i, alg in enumerate(cc.ALG_NAMES.values()):
        ax.bar(x + i * w, [next(b[4] for b in baris if b[0] == n and b[1] == alg) for n in nama],
               w, label=alg, color=WARNA[alg])
    ax.axhline(8, ls="--", c="#3fb950", lw=1, label="maksimum teoretis 8 bit/byte")
    ax.set_xticks(x); ax.set_xticklabels([n[:16] for n in nama], fontsize=7.5, rotation=12)
    ax.set_ylabel("entropi Shannon (bit/byte)"); ax.set_ylim(0, 9); ax.legend(fontsize=7.5, ncol=2)
    ax.set_title("Entropi plainteks vs cipherteks", fontsize=10)
    fig.tight_layout(); fig.savefig(os.path.join(DIR_GRAFIK, "t4_entropi.png")); plt.close(fig)


# ===========================================================================
# T5 — Perbandingan mode, termasuk demonstrasi kelemahan ECB
# ===========================================================================
def t5_ecb_vs_aman():
    """Enkripsi citra dengan ECB dan GCM lalu tampilkan berdampingan.
    Pola gambar masih terbaca pada ECB — inilah alasan ECB dilarang."""
    print("\n[T5] Demonstrasi ECB vs mode aman")
    from PIL import Image
    img = Image.open(os.path.join(DIR_UJI, "06_citra.png")).convert("RGB")
    w, h = img.size
    mentah = img.tobytes()
    kunci = os.urandom(32)

    p = sympad.PKCS7(128).padder()
    ecb = Cipher(algorithms.AES(kunci), modes.ECB()).encryptor()
    c_ecb = ecb.update(p.update(mentah) + p.finalize()) + ecb.finalize()
    c_gcm = cc._aead(cc.ALG_AES_GCM, kunci).encrypt(os.urandom(12), mentah, b"")

    def ke_citra(b):
        return Image.frombytes("RGB", (w, h), b[:w * h * 3])

    fig, ax = plt.subplots(1, 3, figsize=(9.5, 3.9))
    for a, (im, judul) in zip(ax, [
            (img, "Citra asli"),
            (ke_citra(c_ecb), f"AES-256-ECB\nH={_entropi(c_ecb):.3f} — pola masih terbaca"),
            (ke_citra(c_gcm), f"AES-256-GCM\nH={_entropi(c_gcm):.3f} — tampak acak")]):
        a.imshow(im); a.set_title(judul, fontsize=9, pad=8); a.axis("off"); a.grid(False)
    fig.tight_layout(); fig.savefig(os.path.join(DIR_GRAFIK, "t5_ecb_vs_gcm.png")); plt.close(fig)
    print(f"      entropi ECB {_entropi(c_ecb):.4f} vs GCM {_entropi(c_gcm):.4f} bit/byte")

    return [
        ["AES-256-GCM", "AEAD (CTR + GHASH)", 32, 12, 16, "Ya", "Ya",
         "Dipercepat AES-NI di CPU modern", "FITUR UTAMA"],
        ["ChaCha20-Poly1305", "AEAD (stream ARX)", 32, 12, 16, "Ya", "Ya",
         "Cepat tanpa AES-NI, waktu konstan", "FITUR UTAMA"],
        ["AES-256-CBC", "Blok, rantai", 32, 16, 0, "Ya", "Tidak",
         "Rentan padding-oracle bila tanpa MAC", "PEMBANDING"],
        ["AES-256-ECB", "Blok, tanpa rantai", 32, 0, 0, "Tidak", "Tidak",
         "Pola plainteks bocor — DILARANG", "PEMBANDING"],
    ]


# ===========================================================================
# Penulisan XLSX
# ===========================================================================
JUDUL = Font(bold=True, color="FFFFFF")
ISI_JUDUL = PatternFill("solid", fgColor="1F3864")


def tulis_sheet(wb, nama, header, baris, lebar=None):
    ws = wb.create_sheet(nama)
    ws.append(header)
    for c in ws[1]:
        c.font = JUDUL; c.fill = ISI_JUDUL; c.alignment = Alignment(wrap_text=True, vertical="center")
    for b in baris:
        ws.append(b)
    for i, _ in enumerate(header, 1):
        ws.column_dimensions[get_column_letter(i)].width = (lebar or [18] * len(header))[i - 1]
    ws.freeze_panes = "A2"
    return ws


def main():
    if not os.path.isdir(DIR_UJI) or len(os.listdir(DIR_UJI)) < 5:
        sys.exit("Data uji belum ada. Jalankan dulu: python data/buat_data_uji.py")

    t1 = t1_kebenaran()
    t2, kdf = t2_waktu()
    t3 = t3_avalanche()
    t4 = t4_entropi()
    t5 = t5_ecb_vs_aman()

    wb = Workbook(); wb.remove(wb.active)

    ws = wb.create_sheet("Ringkasan")
    for r in [["HASIL PENGUJIAN — Brankas Berkas Pribadi Terenkripsi"], [],
              ["Kasus uji kebenaran", len(t1)],
              ["Kasus lulus penuh",
               sum(1 for b in t1 if b[5] == "BERHASIL" and b[6] == "DITOLAK" and b[7] == "DITOLAK")],
              ["Berkas uji", len(berkas_uji())],
              ["Algoritma utama", " / ".join(cc.ALG_NAMES.values())],
              ["Pembanding (tidak dipakai sbg fitur utama)", "AES-256-CBC, AES-256-ECB"],
              ["Percobaan avalanche per skema", 100],
              [], ["Lembar", "Isi"],
              ["T1 Kebenaran", "Roundtrip + penolakan sandi salah & data diubah"],
              ["T2 Waktu", "Waktu enkripsi/dekripsi 1 KB / 1 MB / 10 MB + biaya KDF"],
              ["T3 Avalanche", "Perubahan bit cipherteks akibat 1 bit plainteks/kunci"],
              ["T4 Entropi", "Entropi Shannon, chi-square, histogram"],
              ["T5 Perbandingan", "Perbandingan algoritma & mode"]]:
        ws.append(r)
    ws["A1"].font = Font(bold=True, size=13)
    ws.column_dimensions["A"].width = 42; ws.column_dimensions["B"].width = 58

    tulis_sheet(wb, "T1 Kebenaran",
                ["Berkas", "Algoritma", "Ukuran asli (byte)", "Ukuran kontainer (byte)",
                 "Overhead (byte)", "Roundtrip", "Sandi salah", "1 byte diubah"], t1,
                [26, 20, 16, 18, 13, 13, 13, 14])
    tulis_sheet(wb, "T2 Waktu",
                ["Algoritma", "Ukuran", "Byte", "Enkripsi (ms)", "Dekripsi (ms)",
                 "Std dev (ms)", "Throughput (MB/s)"], t2, [22, 10, 14, 14, 14, 13, 17])
    ws = wb["T2 Waktu"]; ws.append([]); ws.append(["BIAYA KDF (sekali per operasi)"])
    ws.append(["KDF", "Waktu (ms)", "Status"])
    for k in kdf:
        ws.append(k)
    tulis_sheet(wb, "T3 Avalanche",
                ["Algoritma/Mode", "Skenario", "Badan cipherteks (%)", "Std dev (%)",
                 "Auth tag (%)", "Acuan"], t3, [22, 24, 20, 13, 14, 12])
    tulis_sheet(wb, "T4 Entropi",
                ["Berkas", "Algoritma", "Byte diuji", "H plainteks", "H cipherteks",
                 "Chi2 plainteks", "Chi2 cipherteks", "Kesimpulan"], t4,
                [24, 20, 13, 13, 14, 15, 16, 15])
    tulis_sheet(wb, "T5 Perbandingan",
                ["Algoritma/Mode", "Jenis", "Kunci (byte)", "IV/Nonce (byte)", "Tag (byte)",
                 "Kerahasiaan", "Keutuhan", "Catatan", "Status di aplikasi"], t5,
                [22, 22, 12, 15, 12, 13, 12, 40, 18])

    out = os.path.join(DIR_HASIL, "hasil_pengujian.xlsx")
    wb.save(out)
    print(f"\nSelesai.\n  XLSX   : {out}\n  Grafik : {DIR_GRAFIK}")
    for g in sorted(os.listdir(DIR_GRAFIK)):
        print("           ", g)


if __name__ == "__main__":
    main()
